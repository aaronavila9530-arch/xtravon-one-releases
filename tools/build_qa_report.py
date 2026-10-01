from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUT = ROOT / "Entregables"
OUT.mkdir(exist_ok=True)

DOCX_PATH = OUT / "XTRAVON_ONE_Reporte_Tecnico_QA.docx"
PDF_PATH = OUT / "XTRAVON_ONE_Reporte_Tecnico_QA.pdf"
IMG_DIR = OUT / "qa_report_assets"
IMG_DIR.mkdir(exist_ok=True)

NAVY = "081324"
ATLANTIC = "0B1E3A"
CYAN = "00E5FF"
COMMAND = "0066FF"
MINT = "2BD4A7"
AMBER = "FFB21F"
RISK = "FF4F7A"
ICE = "BFFBFF"
WHITE = "F5F7FA"
STEEL = "C7CDD6"
INK = "111827"
MUTED = "4B5563"


def font(size: int, bold: bool = False):
    candidates = [
        "C:/Windows/Fonts/aptos.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibri.ttf",
    ]
    bold_candidates = [
        "C:/Windows/Fonts/aptosbd.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/calibrib.ttf",
    ]
    for path in (bold_candidates if bold else candidates):
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def set_cell_fill(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "17324B", size: str = "8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_run(run, size=10, color=WHITE, bold=False):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_text(paragraph, text, size=10, color=WHITE, bold=False):
    run = paragraph.add_run(text)
    style_run(run, size=size, color=color, bold=bold)
    return run


def add_heading(doc: Document, title: str, subtitle: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    add_text(p, title, size=18, color=CYAN, bold=True)
    if subtitle:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(8)
        add_text(sp, subtitle, size=9, color=MUTED)


def make_cover() -> Path:
    w, h = 1800, 1050
    img = Image.new("RGB", (w, h), "#" + NAVY)
    draw = ImageDraw.Draw(img)

    for i in range(0, w, 70):
        draw.line((i, 0, i - 350, h), fill="#0C223B", width=1)
    for y in range(0, h, 70):
        draw.line((0, y, w, y), fill="#0C223B", width=1)

    logo = Image.open(ASSETS / "xtravon_splash_source.png").convert("RGBA")
    logo.thumbnail((520, 520))
    img.paste(logo, (70, 90), logo)

    draw.text((650, 120), "XTRAVON ONE", font=font(72, True), fill="#" + WHITE)
    draw.text((655, 205), "REPORTE TECNICO QA", font=font(46, True), fill="#" + CYAN)
    draw.text(
        (655, 270),
        "Revision integral de backend, escritorio, app mobile y contrato DB/API",
        font=font(28),
        fill="#" + STEEL,
    )

    cards = [
        ("Backend", "Compila e importa", MINT),
        ("OpenAPI", "125 paths generados", CYAN),
        ("Mobile", "Bundle Android OK", COMMAND),
        ("DB/API", "Contrato alineado", MINT),
        ("Riesgos", "4 hallazgos P1", RISK),
        ("Dependencias", "13 moderadas", AMBER),
    ]
    x0, y0 = 655, 390
    for idx, (title, value, color) in enumerate(cards):
        x = x0 + (idx % 3) * 355
        y = y0 + (idx // 3) * 155
        draw.rounded_rectangle((x, y, x + 315, y + 110), radius=14, fill="#0B1E3A", outline="#" + color, width=3)
        draw.rectangle((x, y, x + 315, y + 12), fill="#" + color)
        draw.text((x + 22, y + 30), title, font=font(25, True), fill="#" + WHITE)
        draw.text((x + 22, y + 68), value, font=font(19), fill="#" + STEEL)

    draw.line((70, 875, 1730, 875), fill="#" + CYAN, width=3)
    draw.text((70, 910), "QORA SYSTEMS - XTRAVON ONE | GRAIN CONTROL", font=font(24, True), fill="#" + WHITE)
    draw.text((70, 955), f"Fecha de revision: {datetime.now().strftime('%d/%m/%Y')}", font=font(20), fill="#" + STEEL)
    draw.text((1330, 955), "Documento tecnico interno", font=font(20), fill="#" + STEEL)

    path = IMG_DIR / "qa_cover.png"
    img.save(path)
    return path


def make_risk_matrix() -> Path:
    w, h = 1600, 850
    img = Image.new("RGB", (w, h), "#" + NAVY)
    draw = ImageDraw.Draw(img)
    draw.text((55, 40), "Mapa QA de Riesgo", font=font(48, True), fill="#" + WHITE)
    draw.text((55, 105), "Impacto operativo vs. probabilidad tecnica", font=font(24), fill="#" + STEEL)
    left, top, cell = 180, 190, 155
    labels = ["Baja", "Media", "Alta", "Critica"]
    for i in range(4):
        draw.text((left + i * cell + 48, top - 45), labels[i], font=font(18, True), fill="#" + STEEL)
        draw.text((35, top + i * cell + 55), labels[3 - i], font=font(18, True), fill="#" + STEEL)
    for r in range(4):
        for c in range(4):
            score = (4 - r) + c
            color = MINT if score <= 4 else AMBER if score <= 6 else RISK
            draw.rectangle((left + c * cell, top + r * cell, left + (c + 1) * cell - 6, top + (r + 1) * cell - 6), fill="#" + color)
    draw.text((465, 780), "Probabilidad", font=font(23, True), fill="#" + CYAN)
    draw.text((18, 500), "Impacto", font=font(23, True), fill="#" + CYAN)
    findings = [
        ("P1", "QR_SECRET por defecto", 3, 0),
        ("P1", "CORS abierto", 2, 1),
        ("P1", "Expiracion QR 24h", 3, 1),
        ("P1", "node_modules en Git", 1, 0),
        ("P2", "DDL en routers", 2, 2),
        ("P2", "IA sin circuit breaker", 2, 2),
    ]
    for sev, text, c, r in findings:
        x = left + c * cell + 18
        y = top + r * cell + 36
        draw.rounded_rectangle((x, y, x + 118, y + 52), radius=8, fill="#" + NAVY, outline="#" + WHITE)
        draw.text((x + 10, y + 8), sev, font=font(17, True), fill="#" + WHITE)
    x0 = 900
    draw.text((x0, 205), "Hallazgos principales", font=font(32, True), fill="#" + CYAN)
    y = 270
    for sev, text, _, _ in findings:
        color = RISK if sev == "P1" else AMBER
        draw.rounded_rectangle((x0, y, x0 + 88, y + 42), radius=8, fill="#" + color)
        draw.text((x0 + 22, y + 8), sev, font=font(18, True), fill="#" + NAVY)
        draw.text((x0 + 110, y + 8), text, font=font(22, True), fill="#" + WHITE)
        y += 68
    path = IMG_DIR / "qa_risk_matrix.png"
    img.save(path)
    return path


def add_table(doc, headers, rows, widths=None, header_fill=CYAN):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, head in enumerate(headers):
        set_cell_fill(hdr[i], header_fill)
        set_cell_border(hdr[i], "0B1E3A")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, head, size=8, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_fill(cells[i], NAVY if i != 0 else ATLANTIC)
            set_cell_border(cells[i], "17324B")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            add_text(p, str(value), size=8, color=WHITE if i != 0 else CYAN, bold=i == 0)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def build_docx():
    cover = make_cover()
    matrix = make_risk_matrix()
    flowchart = OUT / "XTRAVON_ONE_Flowchart_End_to_End_ES.jpg"

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(cover), width=Inches(7.35))
    doc.add_page_break()

    add_heading(doc, "Resumen ejecutivo", "Quality Assurance tecnico del ecosistema XTRAVON ONE.")
    p = doc.add_paragraph()
    add_text(
        p,
        "El sistema paso las pruebas base de compilacion, importacion FastAPI, generacion OpenAPI, contrato DB/API y bundle Android. "
        "Sin embargo, hay riesgos de produccion que deben corregirse antes de considerar el producto endurecido para operacion portuaria continua.",
        size=10,
        color=INK,
    )

    add_table(
        doc,
        ["Area", "Resultado", "Lectura QA"],
        [
            ["Backend Python", "OK", "Compila sin errores de sintaxis."],
            ["FastAPI/OpenAPI", "OK", "146 rutas importadas; 125 paths OpenAPI generados."],
            ["Contrato DB/API", "OK", "Tablas principales y PUT endpoints alineados."],
            ["Mobile Android", "OK", "Expo export Android completo; bundle generado."],
            ["Dependencias mobile", "Atencion", "13 vulnerabilidades moderadas; fix automatico implica cambio mayor."],
            ["Release hygiene", "Riesgo alto", "node_modules aparece rastreado/modificado en Git."],
        ],
        widths=[1.7, 1.2, 4.5],
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(matrix), width=Inches(7.2))
    doc.add_page_break()

    add_heading(doc, "Hallazgos priorizados", "Ordenados por severidad operativa y probabilidad tecnica.")
    findings = [
        (
            "P1",
            "Mobile release hygiene",
            "mobile/node_modules aparece rastreado/modificado aunque .gitignore lo ignora.",
            "Builds pesados, commits contaminados y diferencias dificiles de auditar.",
            "Remover node_modules del indice Git y dejar package-lock como fuente de verdad.",
        ),
        (
            "P1",
            "QR_SECRET inseguro por defecto",
            "backend/database.py define secreto por defecto; qr_security solo falla si ENVIRONMENT=production.",
            "Riesgo de QR falsificables si Railway no tiene variables correctas.",
            "Exigir QR_SECRET real y ENVIRONMENT=production en arranque.",
        ),
        (
            "P1",
            "CORS abierto",
            "backend/main.py usa allow_origins=['*'] con allow_credentials=True.",
            "Exposicion innecesaria del backend productivo.",
            "Configurar allowlist por entorno y origen esperado.",
        ),
        (
            "P1",
            "Expiracion QR max 24h",
            "despacho_viajes.py limita expira_minutos a 1440.",
            "No alinea con operaciones de varios dias y QR precargados.",
            "Separar guia precargada de QR activo por ciclo vigente.",
        ),
        (
            "P2",
            "DDL dentro de routers",
            "base_operaciones, despacho, issue_log y aprobaciones ejecutan asegurar_estructura en runtime.",
            "Primer request lento, locks o fallas intermitentes.",
            "Mover a migraciones/preflight de deploy.",
        ),
        (
            "P2",
            "PORTIA/IA en request sin circuit breaker",
            "ai_assistant.py consume APIs externas y OpenAI en linea.",
            "Lag, timeout o 502 si un proveedor tarda.",
            "Cache, timeouts cortos, fallback y colas de analisis.",
        ),
        (
            "P2",
            "API hardcoded en app",
            "mobile/src/config.js apunta directo a Railway productivo.",
            "Dificulta staging, pruebas y soporte de handheld/celular.",
            "Usar EXPO_PUBLIC_API_BASE y perfiles EAS por entorno.",
        ),
        (
            "P2",
            "Offline QR cache sensible",
            "ScanScreen usa token/hash en cache offline.",
            "Necesario para operar sin red, pero debe protegerse.",
            "SecureStore/SQLite protegida, binding por dispositivo y token por viaje.",
        ),
    ]
    add_table(doc, ["Sev.", "Hallazgo", "Evidencia", "Impacto", "Recomendacion"], findings, widths=[0.55, 1.35, 2.1, 1.6, 1.8])

    add_heading(doc, "Evidencias tecnicas", "Comandos ejecutados y resultado observado.")
    checks = [
        ["python -m compileall", "PASS", "backend, frontend, main.py y tools compilan."],
        ["FastAPI import", "PASS", "main.app importa; 146 rutas registradas."],
        ["OpenAPI generation", "PASS", "125 paths generados sin excepcion."],
        ["Duplicate route scan", "PASS", "No se detectaron path+method duplicados."],
        ["DB contract audit", "PASS", "Tablas y PUT principales alineados."],
        ["Expo Android export", "PASS", "Bundle Android generado correctamente."],
        ["npm audit", "ATTENTION", "13 vulnerabilidades moderadas; fix force cambia Expo."],
    ]
    add_table(doc, ["Check", "Estado", "Detalle"], checks, widths=[2.0, 1.0, 4.3], header_fill=COMMAND)

    add_heading(doc, "Arquitectura y flujo operativo", "El QA valida el sistema como una operacion end-to-end, no como pantallas aisladas.")
    if flowchart.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(flowchart), width=Inches(7.25))
    else:
        p = doc.add_paragraph()
        add_text(p, "Flowchart no encontrado en Entregables.", color=AMBER, bold=True)

    add_heading(doc, "Plan de correccion recomendado", "Secuencia sugerida para corregir sin romper operacion.")
    plan_rows = [
        ["1", "Limpiar repo mobile", "Remover node_modules del indice, confirmar package-lock, build limpio.", "Alto"],
        ["2", "Endurecer secretos", "QR_SECRET obligatorio, ENVIRONMENT productivo, validacion al arranque.", "Alto"],
        ["3", "Redisenar expiracion QR", "Guia precargada no expira; QR activo expira por ciclo confirmado.", "Alto"],
        ["4", "Migraciones", "Mover asegurar_estructura a scripts de deploy/versionado.", "Medio"],
        ["5", "PORTIA resiliente", "Cache, timeouts, fallback rapido y respuesta puntual.", "Medio"],
        ["6", "Tests smoke", "Rutas criticas, escaneo QR, offline sync, aprobaciones y despacho.", "Medio"],
    ]
    add_table(doc, ["Orden", "Accion", "Criterio de aceptacion", "Prioridad"], plan_rows, widths=[0.65, 1.7, 4.0, 1.0], header_fill=MINT)

    add_heading(doc, "Conclusion QA", "Estado general del producto.")
    p = doc.add_paragraph()
    add_text(
        p,
        "XTRAVON ONE tiene una base funcional amplia y el contrato tecnico principal compila y responde. "
        "El siguiente salto de calidad no es agregar mas pantallas, sino endurecer seguridad, release management, "
        "migraciones, expiracion QR y resiliencia offline/IA. Corregidos esos puntos, el producto queda mucho mas cerca "
        "de una herramienta comercial robusta para operaciones portuarias de graneles.",
        size=10,
        color=INK,
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_text(footer, "QORA SYSTEMS - XTRAVON ONE | Reporte QA", size=8, color=STEEL)

    # Background-like dark page effect through paragraph/table fills is not native in DOCX;
    # keep all content containers dark to preserve the XTRAVON visual language.
    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_docx()
    print(path)
