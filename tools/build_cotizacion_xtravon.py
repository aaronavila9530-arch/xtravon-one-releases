from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "Entregables"
OUT_FILE = OUT_DIR / "XTRAVON_ONE_Cotizacion_Comercial.docx"
LOGO = ROOT / "assets" / "XTRAVON_seal_round_transparent.png"
SPLASH = ROOT / "assets" / "xtravon_splash_source.png"

NAVY = "081324"
ATLANTIC = "0B1E3A"
CYAN = "00E5FF"
ICE = "6FFBFF"
BLUE = "0066FF"
STEEL = "1B1F26"
SILVER = "C7CDD6"
WHITE = "F5F7FA"
GREEN = "24D9A7"
AMBER = "FFB020"
RED = "FF4D6D"
BLACK = "111827"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D7DCE2", size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_borders(table, color: str = "D7DCE2") -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, color)


def add_run(paragraph, text: str, bold: bool = False, color: str = BLACK, size: int = 10):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(size)
    return run


def add_heading(doc: Document, title: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ATLANTIC if level == 1 else BLACK)
    run.font.size = Pt(17 if level == 1 else 13)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        add_run(p, bold_prefix, bold=True)
        add_run(p, text[len(bold_prefix):])
    else:
        add_run(p, text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        add_run(p, item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        add_run(p, item)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], ATLANTIC)
        set_cell_border(hdr[i], ATLANTIC)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, h, bold=True, color=WHITE, size=9)
        if widths:
            hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            add_run(p, str(value), color=BLACK, size=9)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if widths:
                cells[i].width = Inches(widths[i])
    set_table_borders(table)
    doc.add_paragraph()


def add_kpi_band(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [
        ("Trazabilidad", "5 puntos de control"),
        ("Operación offline", "Handheld con cola local"),
        ("Control de cuota", "Cliente, producto y buque"),
        ("Auditoría", "Usuario, fecha, hora y acción"),
    ]
    for i, (label, value) in enumerate(labels):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, CYAN)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, f"{label}\n", bold=True, color=CYAN, size=10)
        add_run(p, value, bold=True, color=WHITE, size=11)
    doc.add_paragraph()


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, NAVY)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
    left, right = table.rows[0].cells
    p = left.paragraphs[0]
    add_run(p, "XTRAVON ONE\n", bold=True, color=WHITE, size=24)
    add_run(p, "GRAIN CONTROL\n", bold=True, color=CYAN, size=17)
    add_run(p, "Cotización técnica y comercial\n", bold=True, color=WHITE, size=13)
    add_run(p, "Sistema integral para control de operaciones graneleras, trazabilidad QR, despacho, SOF, reportes y continuidad offline.", color=SILVER, size=10)
    if LOGO.exists():
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.add_run().add_picture(str(LOGO), width=Inches(1.55))
    doc.add_paragraph()

    if SPLASH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SPLASH), width=Inches(5.75))
        p.paragraph_format.space_after = Pt(4)

    add_kpi_band(doc)

    meta = doc.add_table(rows=5, cols=2)
    rows = [
        ("Preparado para", "Cliente / operación granelera"),
        ("Preparado por", "QORA SYSTEMS - Alianza con MSL Marine Surveyors and Logistics Group"),
        ("Fecha", date.today().strftime("%d/%m/%Y")),
        ("Versión", "Propuesta editable según requerimientos del cliente"),
        ("Validez referencial", "30 días, sujeto a alcance final, cantidad de handhelds y condiciones contractuales"),
    ]
    for idx, (k, v) in enumerate(rows):
        meta.rows[idx].cells[0].text = k
        meta.rows[idx].cells[1].text = v
        set_cell_shading(meta.rows[idx].cells[0], "E8F4FA")
        for c in meta.rows[idx].cells:
            set_cell_border(c, "C9D5DD")
            for paragraph in c.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string(BLACK)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_page_break()


def add_header_footer(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "XTRAVON ONE | GRAIN CONTROL"
        header.runs[0].font.color.rgb = RGBColor.from_string(ATLANTIC)
        header.runs[0].font.size = Pt(9)
        footer = section.footer.paragraphs[0]
        footer.text = "Cotización editable - QORA SYSTEMS / MSL Marine Surveyors and Logistics Group"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.color.rgb = RGBColor.from_string("6B7280")
        footer.runs[0].font.size = Pt(8)


def build_doc() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor.from_string(BLACK)

    add_cover(doc)
    add_header_footer(doc)

    add_heading(doc, "1. Resumen ejecutivo")
    add_body(
        doc,
        "XTRAVON ONE | GRAIN CONTROL es una plataforma operativa para controlar la descarga, despacho, trazabilidad y auditoría de operaciones graneleras desde la apertura del buque hasta el recibo final en planta o bodega. El sistema integra escritorio, backend, handheld de patio, aplicación móvil de chofer y aplicación móvil de supervisor.",
    )
    add_body(
        doc,
        "El objetivo principal es reducir pérdidas, errores documentales, duplicidad de guías, robo o uso indebido de QR, diferencias de peso y demoras no documentadas, convirtiendo la operación diaria en información accionable para despacho, patio, gerencia y auditoría.",
    )
    add_bullets(
        doc,
        [
            "Controla cada viaje por guía, chofer, placa, cliente, producto, bodega y operación.",
            "Permite operación offline en handheld para no detener patio si falla la señal o el backend.",
            "Genera trazabilidad por QR seguro, firma HMAC, estados operativos, auditoría y SOF.",
            "Centraliza decisiones en despacho: asignar, reasignar, bloquear, liberar y cerrar viajes.",
            "Muestra KPIs, cuotas, descargado real, pendiente de descarga, productividad, informes y alertas.",
        ],
    )

    add_heading(doc, "2. Valor agregado para el cliente")
    add_table(
        doc,
        ["Dolor operativo", "Cómo lo resuelve XTRAVON ONE", "Impacto esperado"],
        [
            ["Papel, lápiz y documentos dispersos", "Digitaliza guías, escaneos, SOF, reportes y auditoría.", "Menos reprocesos y mejor evidencia documental."],
            ["Riesgo de QR falsos o duplicados", "QR firmado, validado contra backend/cache y bloqueado al completar ciclo.", "Mayor control de ingreso y salida."],
            ["Asignaciones rígidas a choferes", "Despacho controla asignación/reasignación y app chofer muestra un QR vigente por ciclo.", "Flexibilidad sin perder seguridad."],
            ["Sin señal en puerto", "Handheld guarda escaneos y SOF localmente y sincroniza cuando vuelve la conexión.", "La operación no se detiene por caída de red."],
            ["Diferencias de carga y cuotas", "Compara cuota vs descargado real por cliente, producto, buque y bodega.", "Evita sobrecuota y mejora liquidación."],
            ["Demoras no documentadas", "SOF captura eventos, fechas, horas, bodega, subcategoría y evidencia operativa.", "Base para reclamos, auditoría y mejora de procesos."],
        ],
        [1.45, 2.9, 2.2],
    )

    add_heading(doc, "3. Alcance funcional incluido")
    add_table(
        doc,
        ["Componente", "Funciones principales"],
        [
            ["ERP Desktop Supervisor", "Apertura de buques, cuotas por cliente, carga inicial de boletas, despacho de viajes, aprobaciones extraordinarias, SOF, centro ejecutivo, informes, roles/permisos, ayuda Q&A y P.O.R.T.I.A."],
            ["Backend / API / Base de datos", "FastAPI, PostgreSQL, validación QR, HMAC, RBAC, auditoría, sincronización offline, reportes, releases y endpoints para escritorio/app/handheld."],
            ["Handheld operador patio", "Lectura SE4710 o cámara de respaldo, captura de escaneos, pesos, ficha, tolva, marchamos, SOF offline y sincronización automática."],
            ["App chofer", "Login por usuario, QR vigente, viajes asignados/pendientes, continuidad controlada, acumulados propios y solicitud de nuevos viajes."],
            ["App supervisor", "Centro ejecutivo móvil, despacho, informes, aprobaciones, SOF, PORTIA y seguimiento operativo alineado con escritorio."],
            ["PORTIA", "Asistente operativo para consultas de riesgo, clima/mar, puertos, operación, cuotas, SOF, duración estimada y apoyo gerencial."],
        ],
        [1.75, 4.85],
    )

    add_heading(doc, "4. Flujo propuesto con 5 escaneos")
    add_body(doc, "La propuesta considera cinco puntos de control, configurables según operación, puerto, terminal, planta y nivel de control requerido.")
    add_table(
        doc,
        ["Escaneo", "Punto de control", "Captura sugerida", "Resultado operativo"],
        [
            ["1", "Báscula Puerto ingreso", "QR, guía, chofer, placa, ficha, peso vacío, fecha/hora y usuario operador.", "Guía pasa a ingreso / en puerto."],
            ["2", "Tolva", "Número de tolva, bodega, producto, evento SOF si aplica y hora exacta.", "Guía pasa a carga / tolva."],
            ["3", "Báscula Puerto salida", "Peso lleno, marchamos múltiples, hora salida, validación de peso neto y bloqueo del QR de puerto.", "Guía completa ciclo portuario."],
            ["4", "Báscula Planta recibo", "Peso de recibo, diferencia contra peso puerto, evidencia o comentario si aplica.", "Valida recepción en destino."],
            ["5", "Bodega recibo", "Confirmación final de descarga, bodega destino, responsable y cierre documental.", "Cierra trazabilidad end-to-end."],
        ],
        [0.65, 1.45, 2.85, 1.75],
    )

    add_heading(doc, "5. Modelo operativo por rol")
    add_table(
        doc,
        ["Rol", "Responsabilidad dentro del sistema"],
        [
            ["Supervisor / Master", "Configura operación, productos, bodegas, cuotas, roles, cierres, reportes y aprobaciones extraordinarias."],
            ["Despacho", "Asigna/reasigna guías, controla choferes disponibles, bloqueos, QR activos y solicitudes de viaje."],
            ["Operador patio", "Escanea QR, registra pesos, tolva, marchamos, SOF y trabaja offline cuando sea necesario."],
            ["Chofer", "Visualiza únicamente su QR vigente, viaje actual, acumulados propios y continuidad autorizada."],
            ["Gerencia / Cliente", "Consulta avance, cuotas, productividad, reportes e indicadores ejecutivos."],
        ],
        [1.6, 5.0],
    )

    add_heading(doc, "6. Cotización económica")
    add_body(doc, "Los valores siguientes son referenciales y editables según alcance final, cantidad de usuarios, cantidad de handhelds, integraciones y necesidades particulares del cliente.")
    add_table(
        doc,
        ["Concepto", "Valor", "Periodicidad", "Incluye"],
        [
            ["Implementación, capacitación y desarrollo inicial", "USD 3,000", "Pago único", "Levantamiento operativo, configuración base, capacitación inicial, ajustes necesarios de arranque y acompañamiento de puesta en marcha."],
            ["Fee plataforma XTRAVON ONE", "USD 3,000", "Mensual", "Uso del sistema, backend, escritorio, apps, reportes, soporte funcional base, mantenimiento evolutivo menor y hosting/licenciamiento base de la plataforma."],
            ["Renting handheld", "USD 50 por equipo", "Mensual", "Uso mensual de cada handheld asignado para operación de patio."],
            ["Personal operativo", "USD 13 por persona/hora", "Según uso", "Tarifa preferencial al adquirir el sistema. Tarifa regular: USD 15 por persona/hora."],
            ["Reposición por daño/pérdida de handheld", "USD 500 por equipo", "Cuando aplique", "Cobro completo si el equipo es dañado, extraviado, manipulado indebidamente o devuelto inutilizable."],
        ],
        [1.8, 1.1, 1.05, 2.75],
    )

    add_heading(doc, "7. Ejemplo de cálculo mensual")
    add_table(
        doc,
        ["Escenario", "Cálculo", "Total mensual estimado"],
        [
            ["Plataforma base", "Fee mensual XTRAVON ONE", "USD 3,000"],
            ["4 handhelds", "4 x USD 50", "USD 200"],
            ["8 handhelds", "8 x USD 50", "USD 400"],
            ["Personal", "Horas reales x USD 13 por persona/hora", "Variable"],
        ],
        [1.7, 3.1, 1.8],
    )

    add_heading(doc, "8. Condiciones de uso y responsabilidades")
    add_bullets(
        doc,
        [
            "La propuesta es editable y puede ajustarse al flujo real del cliente, cantidad de escaneos, integraciones, formatos de informe, usuarios y handhelds requeridos.",
            "El cliente debe cuidar los handhelds entregados en renting. Daño físico, pérdida, robo, manipulación indebida o devolución en mal estado genera cobro de reposición de USD 500 por equipo.",
            "Los QR, usuarios y permisos son personales y operativos. El cliente debe evitar compartir accesos o permitir uso fuera de los roles definidos.",
            "La operación offline permite continuidad en patio; sin embargo, la sincronización final depende de recuperar conexión a red o backend.",
            "El envío automático por correo o WhatsApp requiere configuración de proveedores externos válidos, por ejemplo SMTP/SendGrid y WhatsApp Business/Twilio/Meta. Costos de terceros no incluidos salvo acuerdo expreso.",
            "Cambios mayores, integraciones con ERP externo, básculas, cámaras, GPS comercial, OCR documental o portales adicionales pueden cotizarse por separado.",
            "Los reportes, alertas, indicadores y recomendaciones ejecutivas son herramientas de apoyo. Las decisiones finales corresponden al cliente y a sus responsables operativos.",
            "Los impuestos, IVA, viáticos, traslados, hospedajes o gastos fuera de la operación remota/local acordada se cotizan o facturan según corresponda.",
        ],
    )

    add_heading(doc, "9. Entregables de implementación")
    add_numbered(
        doc,
        [
            "Configuración inicial del backend, base de datos, usuarios, roles y permisos.",
            "Instalación/configuración del ERP Desktop para supervisor y administración.",
            "Configuración de app supervisor, app chofer y handheld operador patio.",
            "Plantilla de carga de boletas y guías por operación.",
            "Flujo QR con cinco escaneos operativos.",
            "Centro Ejecutivo con KPIs, avance, cuotas, descargado, pendiente, productividad y riesgos.",
            "Informes descargables según formatos acordados.",
            "Capacitación inicial por rol: supervisor, despacho, patio, chofer y gerencia.",
            "Acompañamiento de salida en vivo y ajustes iniciales razonables.",
        ],
    )

    add_heading(doc, "10. Cronograma referencial")
    add_table(
        doc,
        ["Fase", "Duración estimada", "Resultado"],
        [
            ["Diagnóstico y parametrización", "3 a 5 días hábiles", "Flujo validado, roles definidos, plantillas y operación base."],
            ["Configuración técnica", "5 a 10 días hábiles", "Backend, escritorio, apps, handhelds y usuarios operativos."],
            ["Pruebas piloto", "3 a 7 días hábiles", "Prueba con buque/operación simulada o real controlada."],
            ["Capacitación y salida en vivo", "1 a 3 días hábiles", "Usuarios entrenados y operación en producción."],
            ["Estabilización", "2 a 4 semanas", "Ajustes finos, soporte y mejora de reportes/alertas."],
        ],
        [1.6, 1.45, 3.55],
    )

    add_heading(doc, "11. Supuestos")
    add_bullets(
        doc,
        [
            "El cliente entregará datos maestros mínimos: clientes, productos, choferes, placas, cuotas, bodegas, usuarios y reglas operativas.",
            "La cantidad final de handhelds se definirá antes de iniciar operación.",
            "Las integraciones con básculas, correo, WhatsApp, GPS o sistemas externos dependen de acceso técnico y credenciales del proveedor correspondiente.",
            "El alcance de desarrollo incluido cubre ajustes razonables de implementación. Nuevos módulos mayores se cotizan por separado.",
            "La tarifa mensual no incluye reposición de equipos dañados ni horas de personal operativo.",
        ],
    )

    add_heading(doc, "12. Cierre comercial")
    add_body(
        doc,
        "XTRAVON ONE no es únicamente un sistema de captura: es una capa de control operativo, trazabilidad y gestión de riesgo diseñada para que una operación granelera pueda saber qué se cargó, quién lo movió, dónde se validó, cuánto se descargó, qué falta y qué evento afectó la operación.",
    )
    add_body(
        doc,
        "La propuesta busca iniciar con una base robusta y editable, dejando espacio para adaptar el sistema a las reglas reales del cliente, sin perder seguridad, auditoría ni continuidad operativa.",
    )

    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build_doc()
